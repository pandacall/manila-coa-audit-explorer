"""Read an AAR's Financial Statements (Part I) and Annexes (Part IV) into long-format lines.

Each year's two spreadsheets are read sheet by sheet. A sheet is a statement because its title says
so ("STATEMENT OF CASH FLOWS"), never because of its name or position: sheet names change between
years (2024's Annexes drop the "Annex A - " prefix), sheets change order, and 2023 swaps the Annex
letters of the cash flow and net assets statements. Hidden sheets are working papers and are
ignored. Amounts are the values Excel last cached for the formulas, rounded to the centavo.

The tables share a shape: label cells on the left, whose column shows the indent; value columns
named by a header row (a Fund, "Total", a year, or SCBAA's budget columns); heading rows that
carry no amounts and open a section. A heading row that follows an amount row and finishes its label
("... Disposal of" / "Property, Plant and Equipment") is joined to it instead.

A line is one amount: its AAR year, the statement and where it is printed (Part I or an Annex), its
Fund, section, line item and column. Part I statements are for the City as a whole ("All Funds");
Annexes break the same statements down by Fund.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ALL_FUNDS = "All Funds"
GENERAL_FUND = "General Fund"
SPECIAL_EDUCATION_FUND = "Special Education Fund"
TRUST_FUND = "Trust Fund"
FUNDS = (GENERAL_FUND, SPECIAL_EDUCATION_FUND, TRUST_FUND, ALL_FUNDS)

AMOUNT = "Amount"
PRIOR_YEAR = "Prior-year comparative"
SCBAA_COLUMNS = (
    "Original budget",
    "Final budget",
    "Difference original and final budget",
    "Actual",
    "Difference final budget and actual",
)

STATEMENT_NAMES = {
    "SFPo": "Statement of Financial Position",
    "SFPe": "Statement of Financial Performance",
    "SCNAE": "Statement of Changes in Net Assets/Equity",
    "SCF": "Statement of Cash Flows",
    "SCBAA": "Statement of Comparison of Budget and Actual Amounts",
}
# Matched against a sheet's upper-cased title line.
STATEMENT_TITLES = (
    ("FINANCIAL POSITION", "SFPo"),
    ("FINANCIAL PERFORMANCE", "SFPe"),
    ("CHANGES IN NET ASSETS", "SCNAE"),
    ("CASH FLOW", "SCF"),
    ("COMPARISON OF BUDGET", "SCBAA"),
)

HEADER_SCAN_ROWS = 12
TWO_PLACES = Decimal("0.01")
ENUMERATOR_CELL = re.compile(r"^(?:[A-Za-z]|\d{1,2})\.$")
TOTAL_LABEL = re.compile(r"(?:total|net)\b", re.IGNORECASE)
ENUMERATOR_PREFIX = re.compile(r"^(?:[A-Za-z]|\d{1,2})\.\s+")
ANNEX_LETTER = re.compile(r"^annex\s+([A-Z](?:-\d)?)$", re.IGNORECASE)
YEAR = re.compile(r"^(?:19|20)\d{2}$")
# A label that ends like this goes on in the next row.
OPEN_ENDING = re.compile(
    r"(?:\b(?:of|and|by|in|to|for|the|on|from)|[,(/-]"
    r"|\(Used in\)|Operating, Investing|Other Operating)$",
    re.IGNORECASE,
)
WRAPPED_IN_CELL = re.compile(r"\n\s*$")  # the author broke the line inside the cell


@dataclass
class FinancialLine:
    key: str  # what the model cites; shared by the columns and the fund of one printed row
    source: str  # "Part I" or "Annex A": where the statement is printed
    statement: str  # SFPo, SFPe, SCNAE, SCF or SCBAA
    sheet: str
    row: int  # the spreadsheet row, for tracing a line back to the cell
    fund: str
    section: str  # the headings above the line, outermost first, joined by " > "
    line_item: str
    column: str  # AMOUNT, PRIOR_YEAR, or one of SCBAA_COLUMNS
    period: int  # the year the amount is for
    amount: Decimal
    citation: str

    def to_dict(self) -> dict:
        record = asdict(self)
        record["amount"] = str(self.amount)
        return record


@dataclass
class FinancialRecord:
    aar_year: int
    source_files: list[str]
    lines: list[FinancialLine] = field(default_factory=list)
    ignored_sheets: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "aar_year": self.aar_year,
            "source_files": self.source_files,
            "ignored_sheets": self.ignored_sheets,
            "lines": [line.to_dict() for line in self.lines],
        }


def extract_year(reports_dir: Path, year: int) -> FinancialRecord:
    """Every financial line of one AAR's Financial Statements and Annexes."""
    folder = reports_dir / f"Manila-City-Annual-Audit-Report-{year}"
    part1 = find_workbook(folder, "07-*")
    part4 = find_workbook(folder, "11-*")
    record = FinancialRecord(
        aar_year=year,
        source_files=[path.relative_to(reports_dir).as_posix() for path in (part1, part4)],
    )
    for path in (part1, part4):
        read_workbook(path, year, record)
    return record


def find_workbook(folder: Path, pattern: str) -> Path:
    matches = sorted(p for p in folder.glob(f"**/{pattern}.xlsx") if not p.name.startswith("~$"))
    if len(matches) != 1:
        raise FileNotFoundError(f"expected one {pattern}.xlsx under {folder}: {matches}")
    return matches[0]


def read_workbook(path: Path, year: int, record: FinancialRecord) -> None:
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    for sheet in values.worksheets:
        if sheet.sheet_state != "visible":
            record.ignored_sheets.append({"sheet": sheet.title, "reason": "hidden"})
            continue
        statement = statement_of(sheet)
        if statement is None:
            record.ignored_sheets.append({"sheet": sheet.title, "reason": "not a statement"})
            continue
        lines = read_sheet(sheet, formulas[sheet.title], statement, year)
        add_citations(lines, year)
        record.lines.extend(lines)


def statement_of(sheet: Worksheet) -> str | None:
    for row in sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True):
        for value in row:
            if isinstance(value, str) and value.strip().upper().startswith("STATEMENT OF"):
                title = value.upper()
                return next((code for words, code in STATEMENT_TITLES if words in title), None)
    return None


def annex_letter(sheet: Worksheet) -> str | None:
    for row in sheet.iter_rows(min_row=1, max_row=2, values_only=True):
        for value in row:
            match = isinstance(value, str) and ANNEX_LETTER.match(value.strip())
            if match:
                return match.group(1).upper()
    return None


def fund_of_sheet(title: str) -> str:
    """The Fund a budget-and-actual Annex sheet is for, from its name ("SCBAA-GF", "... SEF")."""
    tokens = set(re.findall(r"[A-Za-z]+", title))
    if "GF" in tokens:
        return GENERAL_FUND
    if "SEF" in tokens:
        return SPECIAL_EDUCATION_FUND
    raise ValueError(f"cannot tell which Fund the sheet {title!r} is for")


class Sections:
    """The headings above the current row, outermost first, each with the column it is indented to.

    A heading closes the sections at its own indent or deeper that already hold lines or deeper
    headings, but not an earlier heading at the same indent that holds neither (SCF's "Cash Flows
    from Operating Activities" and the "Cash Inflows" under it sit in the same column). A
    "Total ..." or "Net ..." line closes the sections it sums.
    """

    def __init__(self) -> None:
        self._open: list[list] = []  # [indent, text, holds lines]

    def open(self, indent: int, text: str) -> None:
        while self._open and (
            self._open[-1][0] > indent or (self._open[-1][0] == indent and self._open[-1][2])
        ):
            self._open.pop()
        if self._open and indent > self._open[-1][0]:
            self._open[-1][2] = True  # it now holds a deeper heading
        self._open.append([indent, text, False])

    def line(self, indent: int, *, total: bool) -> None:
        if total:
            while self._open and self._open[-1][0] >= indent:
                self._open.pop()
        if self._open:
            self._open[-1][2] = True

    def path(self) -> str:
        return " > ".join(text for _, text, _ in self._open)


@dataclass
class ValueColumn:
    index: int  # 1-based
    fund: str
    column: str
    period: int


def read_sheet(
    sheet: Worksheet, formulas: Worksheet, statement: str, year: int
) -> list[FinancialLine]:
    letter = annex_letter(sheet)
    source = f"Annex {letter}" if letter else "Part I"
    columns, first_row, note_column = value_columns(sheet, statement, year, letter is not None)
    label_columns = [c for c in range(1, min(c.index for c in columns)) if c != note_column]
    last_column = max(c.index for c in columns)
    lines: list[FinancialLine] = []
    sections = Sections()
    previous_row: list[FinancialLine] | None = None
    previous_number = 0
    previous_raw = ""
    for number, cells in enumerate(
        sheet.iter_rows(min_row=first_row, max_col=last_column), start=first_row
    ):
        texts = [
            (cell.column, clean(cell.value))
            for cell in cells
            if cell.column in label_columns and isinstance(cell.value, str) and clean(cell.value)
        ]
        texts = [(c, t) for c, t in texts if not ENUMERATOR_CELL.match(t)]
        amounts = {}
        for column in columns:
            cell = cells[column.index - 1]
            if isinstance(cell.value, bool) or not isinstance(cell.value, int | float):
                continue
            amounts[column.index] = cell.value
        for column in columns:
            if column.index not in amounts and cached_missing(
                sheet, formulas, number, column.index
            ):
                raise ValueError(
                    f"{sheet.title} row {number}: a formula has no cached value; open the file in"
                    " Excel and save it"
                )
        if not texts:
            continue
        indent, label = texts[0][0], strip_enumerator(" ".join(t for _, t in texts))
        raw_label = str(next(c.value for c in cells if c.column == texts[-1][0]))
        if not amounts:
            if (
                previous_row
                and previous_number == number - 1
                and continues(previous_row[0].line_item, previous_raw, label)
            ):
                joined = f"{previous_row[0].line_item} {label}"
                for line in previous_row:
                    line.line_item = joined
                continue
            sections.open(indent, label)
            previous_row = None
            continue
        row_lines = [
            FinancialLine(
                key="",
                source=source,
                statement=statement,
                sheet=sheet.title,
                row=number,
                fund=column.fund,
                section=sections.path(),
                line_item=label,
                column=column.column,
                period=column.period,
                amount=pesos(amounts[column.index]),
                citation="",
            )
            for column in columns
            if column.index in amounts
        ]
        lines.extend(row_lines)
        sections.line(indent, total=TOTAL_LABEL.match(label) is not None)
        previous_row, previous_number, previous_raw = row_lines, number, raw_label
    for line in lines:
        line.key = line_key(year, line)
    return lines


def value_columns(
    sheet: Worksheet, statement: str, year: int, annex: bool
) -> tuple[list[ValueColumn], int, int | None]:
    """The value columns named by the header row, the first row of data, and the Note column."""
    rows = list(sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS))
    for position, cells in enumerate(rows):
        named = [(c.column, header_of(c.value)) for c in cells if header_of(c.value)]
        if len(named) < 2:
            continue
        note = next(
            (c.column for c in cells if isinstance(c.value, str) and c.value.strip() == "Note"),
            None,
        )
        if statement == "SCBAA":
            sub_row = rows[position + 1] if position + 1 < len(rows) else []
            columns = budget_columns(named, sub_row, sheet, annex)
            return columns, position + 3, note
        columns = []
        for index, header in named:
            if header[0] == "year":
                if header[1] == year:
                    columns.append(ValueColumn(index, ALL_FUNDS, AMOUNT, year))
                elif header[1] == year - 1:
                    columns.append(ValueColumn(index, ALL_FUNDS, PRIOR_YEAR, year - 1))
                else:
                    raise ValueError(f"{sheet.title}: unexpected year column {header[1]}")
            else:
                columns.append(ValueColumn(index, header[1], AMOUNT, year))
        return columns, position + 2, note
    raise ValueError(f"{sheet.title}: no header row found")


def header_of(value: object) -> tuple[str, object] | None:
    """What a header cell names: ("fund", fund), ("year", 2022) or ("budget", normalised text)."""
    if isinstance(value, int | float) and not isinstance(value, bool):
        value = str(int(value))
    if not isinstance(value, str):
        return None
    text = clean(value)
    lowered = text.lower()
    if YEAR.match(text):
        return ("year", int(text))
    if lowered == "general fund":
        return ("fund", GENERAL_FUND)
    if lowered == "special education fund":
        return ("fund", SPECIAL_EDUCATION_FUND)
    if lowered == "trust fund":
        return ("fund", TRUST_FUND)
    if lowered == "total":
        return ("fund", ALL_FUNDS)
    if lowered.startswith(("budgeted amounts", "difference", "actual amounts")):
        return ("budget", lowered)
    return None


def budget_columns(named, sub_row, sheet: Worksheet, annex: bool) -> list[ValueColumn]:
    """SCBAA's five columns: the budget's Original and Final (a merged header over two columns),
    the difference between them, the Actual, and the difference between Final and Actual."""
    fund = fund_of_sheet(sheet.title) if annex else ALL_FUNDS
    sub = {c.column: clean(c.value).lower() for c in sub_row if isinstance(c.value, str)}
    columns = []
    for index, (_, text) in named:
        if text.startswith("budgeted amounts"):
            by_sub = {sub.get(index): index, sub.get(index + 1): index + 1}
            if set(by_sub) != {"original", "final"}:
                raise ValueError(f"{sheet.title}: cannot find the Original and Final columns")
            columns.append(ValueColumn(by_sub["original"], fund, "Original budget", 0))
            columns.append(ValueColumn(by_sub["final"], fund, "Final budget", 0))
        elif text.startswith("difference original"):
            columns.append(ValueColumn(index, fund, "Difference original and final budget", 0))
        elif text.startswith("actual"):
            columns.append(ValueColumn(index, fund, "Actual", 0))
        elif text.startswith("difference final"):
            columns.append(ValueColumn(index, fund, "Difference final budget and actual", 0))
        else:
            raise ValueError(f"{sheet.title}: unknown column {text!r}")
    return columns


def cached_missing(sheet: Worksheet, formulas: Worksheet, row: int, column: int) -> bool:
    raw = formulas.cell(row=row, column=column).value
    return isinstance(raw, str) and raw.startswith("=") and sheet.cell(row, column).value is None


def continues(previous_label: str, previous_raw: str, label: str) -> bool:
    """Does this amount-less row finish the label of the amount row just above it?"""
    return (
        label[:1].islower()
        or bool(OPEN_ENDING.search(previous_label))
        or bool(WRAPPED_IN_CELL.search(previous_raw))
    )


def clean(value: object) -> str:
    return " ".join(str(value).replace("\xa0", " ").split())


def strip_enumerator(label: str) -> str:
    return ENUMERATOR_PREFIX.sub("", label)


def pesos(value: int | float) -> Decimal:
    return Decimal(str(value)).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)


def line_key(year: int, line: FinancialLine) -> str:
    fund = {GENERAL_FUND: "GF", SPECIAL_EDUCATION_FUND: "SEF", TRUST_FUND: "TF", ALL_FUNDS: "ALL"}
    where = line.source.replace(" ", "")
    return f"{year}-FS-{where}-{line.statement}-{line.row}-{fund[line.fund]}"


def add_citations(lines: list[FinancialLine], year: int) -> None:
    """Cite each line by statement and line item. A line item printed more than once in its sheet
    (every "Total Cash Inflows") gets the headings above it, as few as tell the rows apart."""
    paths: dict[str, dict[int, tuple[str, ...]]] = {}
    for line in lines:
        path = tuple(line.section.split(" > ")) if line.section else ()
        paths.setdefault(line.line_item, {})[line.row] = path
    for line in lines:
        item = line.line_item
        same = paths[item]
        if len(same) > 1:
            depth = next(
                (d for d in range(1, 10) if len({p[-d:] for p in same.values()}) == len(same)),
                10,
            )
            tail = same[line.row][-depth:]
            if tail:
                item = f"{' > '.join(tail)}: {item}"
            if sum(1 for p in same.values() if p[-depth:] == same[line.row][-depth:]) > 1:
                item = f"{item} (row {line.row})"  # printed twice under the same headings
        name = STATEMENT_NAMES[line.statement]
        if line.source == "Part I":
            where = f"Part I, {name}"
        else:
            fund = f" ({line.fund})" if line.statement == "SCBAA" else ""
            where = f"Part IV, {line.source}, {name}{fund}"
        line.citation = f"CY {year} AAR, {where}, {item}"
